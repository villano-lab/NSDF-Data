"""Which Run76 series does NSDF actually hold?  Writes R76/nsdf_r76_series.csv.

Sources compared:
  * the run log, sheet `Run76` of R76/DataSeriesList.xlsx (what was taken);
  * the manifest bundled in the nsdf-cli package (nsdf_dark_matter_cli/r_dataset.csv), which is what
    `nsdf-cli ls` shows. It is frozen at package-release time and never asks the server;
  * with --probe, the live server. A 1-byte ranged GET on a dump's 0000.bin (URL from the same gen-url
    endpoint `nsdf-cli download` uses) returns 206 for a real dump and 404 for an absent one; the gen-url call
    itself answers 200 either way, so its status is not an existence test.

Run under an interpreter with openpyxl and requests, e.g. /opt/anaconda3/bin/python3:
    python nsdf_r76_availability.py --probe
"""
import argparse
import collections
import csv
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl
import requests

REPO = Path(__file__).resolve().parents[2]
DEFAULT_MANIFEST = "/opt/anaconda3/envs/darkmatter_cli_env/lib/python3.10/site-packages/nsdf_dark_matter_cli/r_dataset.csv"
GENURL = "https://intersect.nationalsciencedatafabric.org/nexus/api/v1/darkmatter/gen-url"

ap = argparse.ArgumentParser()
ap.add_argument("--manifest", default=DEFAULT_MANIFEST)
ap.add_argument("--probe", action="store_true", help="also check the live server (about 1000 small requests)")
ap.add_argument("--out", default=str(REPO / "R76" / "nsdf_r76_series.csv"))
args = ap.parse_args()

_UNIT = {"B": 1 / 2**20, "KiB": 1 / 1024, "MiB": 1.0, "GiB": 1024.0}


def mib(s):
    n, u = re.match(r"([\d.]+)\s*([A-Za-z]+)", s).groups()
    return float(n) * _UNIT[u]


# --- run log
# The sheet carries an Excel AutoFilter that HIDES rows (at the time of writing: the 62 DCRC3-triggered series). openpyxl
# returns hidden rows like any other, so every series is read; the filter state is reported below and each series is
# flagged, so a hidden series can never go unnoticed. The workbook itself is never modified.
ws = openpyxl.load_workbook(REPO / "R76" / "DataSeriesList.xlsx", data_only=True)["Run76"]
hdr = [str(c).strip() if c else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
hidden_rows = {i for i, d in ws.row_dimensions.items() if d.hidden}
log = []
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if r and r[0]:
        d = dict(zip(hdr, r))
        d["series"] = str(r[0]).strip()
        d["hidden_by_filter"] = i in hidden_rows
        log.append(d)


def category(d):
    s = (d.get("source / shield") or "").strip()
    return next((n for n in ("PuBe", "Am", "Na-22") if s.startswith(n)), "no source recorded")


# --- bundled manifest: R76-tagged dumps per series
dumps = collections.defaultdict(list)
for m in csv.DictReader(open(args.manifest)):
    s, n = m["filename"].rsplit("_F", 1)
    if m["rseries"] == "R76":
        dumps[s].append((int(n), mib(m["size"])))

# --- local downloads (a folder with no files inside is the stub a failed `nsdf-cli download` leaves behind)
idx = Path.home() / "idx"
local_ok, local_stub = collections.Counter(), collections.Counter()
if idx.is_dir():
    for d in idx.iterdir():
        mm = re.match(r"^(\d{8}_\d{4})_F\d{4}$", d.name)
        if mm and d.is_dir():
            (local_ok if any(p.is_file() for p in d.rglob("*")) else local_stub)[mm.group(1)] += 1

# --- live probe
session = requests.Session()


def exists(mid):
    for attempt in range(3):
        try:
            urls = session.get(GENURL, params={"filename": mid}, timeout=30).json()["urls"]
            b = next(u["url"] for u in urls if u["key"].endswith("0000.bin"))
            r = session.get(b, headers={"Range": "bytes=0-0"}, timeout=30, stream=True)
            code = r.status_code
            r.close()
            if code in (206, 404):
                return code == 206
            raise RuntimeError(f"status {code}")
        except Exception as e:
            last = e
    raise RuntimeError(f"probe failed for {mid}: {last}")


live = {}
if args.probe:
    def check(s):
        if s in dumps:
            ns = sorted(n for n, _ in dumps[s])
            ok = exists(f"{s}_F{ns[0]:04d}") and exists(f"{s}_F{ns[-1]:04d}")
            beyond = exists(f"{s}_F{ns[-1] + 1:04d}")
            return s, {"listed_dumps_verified": ok, "server_has_more_after_last_listed": beyond}
        found = [n for n in range(1, 6) if exists(f"{s}_F{n:04d}")]
        return s, {"unlisted_but_on_server": bool(found)}

    with ThreadPoolExecutor(max_workers=8) as ex:
        live = dict(ex.map(check, [d["series"] for d in log]))

# --- output
fields = ["series", "type", "source_shield", "HV", "trigger", "duration_min", "on_nsdf", "n_dumps", "first_dump",
          "last_dump", "missing_dumps", "size_gib", "live_check", "hidden_by_filter", "local_dumps", "local_empty_stubs"]
with open(args.out, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fields)
    w.writeheader()
    for d in log:
        s = d["series"]
        ns = sorted(n for n, _ in dumps.get(s, []))
        lv = live.get(s, {})
        if not lv:
            note = ""
        elif s in dumps:
            note = "listed dumps found; none beyond" if lv["listed_dumps_verified"] and not lv["server_has_more_after_last_listed"] else "CHECK"
        else:
            note = "CHECK: on server but not listed" if lv["unlisted_but_on_server"] else "absent (F0001-F0005)"
        w.writerow({
            "series": s, "type": d.get("type") or "", "source_shield": (d.get("source / shield") or "").strip(),
            "HV": d.get("HV") if d.get("HV") is not None else "", "trigger": d.get("trigger") or "",
            "duration_min": d.get("duration [min]") if d.get("duration [min]") is not None else "",
            "on_nsdf": bool(ns), "n_dumps": len(ns), "first_dump": ns[0] if ns else "", "last_dump": ns[-1] if ns else "",
            "missing_dumps": (ns[-1] - ns[0] + 1 - len(ns)) if ns else "",
            "size_gib": round(sum(sz for _, sz in dumps[s]) / 1024, 2) if ns else "",
            "live_check": note, "hidden_by_filter": d["hidden_by_filter"], "local_dumps": local_ok.get(s, 0), "local_empty_stubs": local_stub.get(s, 0)})

# --- summary
n_hidden = sum(d["hidden_by_filter"] for d in log)
print(f"sheet filter: range {ws.auto_filter.ref or 'none'}; {len(hidden_rows)} rows hidden, {n_hidden} of them series "
      f"(all still included below); {len(log) - n_hidden} series visible in Excel")
print(f"run log: {len(log)} series ({dict(collections.Counter(category(d) for d in log))})")
by_cat = collections.defaultdict(lambda: [0, 0, 0, 0.0])
for d in log:
    s, c = d["series"], category(d)
    by_cat[c][0] += 1
    if s in dumps:
        by_cat[c][1] += 1
        by_cat[c][2] += len(dumps[s])
        by_cat[c][3] += sum(sz for _, sz in dumps[s]) / 1024
print(f"{'source':<20}{'series':>7}{'on nsdf':>9}{'dumps':>8}{'GiB':>9}")
for c in ("PuBe", "Am", "Na-22", "no source recorded"):
    n, a, k, g = by_cat[c]
    print(f"{c:<20}{n:>7}{a:>9}{k:>8}{g:>9.1f}")
tot = [sum(v[i] for v in by_cat.values()) for i in range(4)]
print(f"{'total':<20}{tot[0]:>7}{tot[1]:>9}{tot[2]:>8}{tot[3]:>9.1f}")
if live:
    print("live probe: listed series verified:", sum(v.get("listed_dumps_verified", False) for v in live.values()),
          "| listed series with more dumps on the server:", sum(v.get("server_has_more_after_last_listed", False) for v in live.values()),
          "| unlisted series found on the server:", sum(v.get("unlisted_but_on_server", False) for v in live.values()))
print("wrote", args.out)
